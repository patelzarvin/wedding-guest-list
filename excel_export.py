import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import get_db, get_wedding_by_id, get_contacts_by_wedding, get_wedding_stats_db

def export_guest_list_excel(output_path=None, base_url="https://wedding-guest-list-8uhm.onrender.com", wedding_id=None, target_side=None):
    if not output_path:
        output_path = os.path.join(os.path.dirname(__file__), "wedding_guest_list.xlsx")
        
    conn = get_db()
    c = conn.cursor()
    
    # 1. Fetch target wedding
    if wedding_id:
        c.execute("SELECT * FROM weddings WHERE id = ?", (wedding_id,))
    else:
        c.execute("SELECT * FROM weddings ORDER BY id DESC LIMIT 1")
    wedding_row = c.fetchone()
    wedding = dict(wedding_row) if wedding_row else {
        "title": "Aisle & Vows Royal Wedding",
        "groom_name": "Groom",
        "bride_name": "Bride",
        "wedding_date": "Upcoming",
        "wedding_city": "Gujarat"
    }
    wid = wedding.get("id", 1)
    
    # 2. Fetch contacts and stats
    all_contacts = get_contacts_by_wedding(wid) if wedding_row else []
    if target_side:
        contacts = [con for con in all_contacts if con.get("side", "").upper() == target_side.upper()]
    else:
        contacts = all_contacts
        
    stats = get_wedding_stats_db(wid) if wedding_row else {
        "total_invited": len(contacts),
        "confirmed": sum(1 for con in contacts if con.get("status") == "CONFIRMED"),
        "pending": sum(1 for con in contacts if con.get("status") != "CONFIRMED"),
        "total_headcount": sum(con.get("total_headcount", 1) for con in contacts)
    }
    
    conn.close()
    
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    
    # Theme 2 Styling Tokens
    header_fill = PatternFill(start_color="182030", end_color="182030", fill_type="solid") # Dark Slate
    header_font = Font(name="Calibri", size=11, bold=True, color="E5C158") # Champagne Gold
    gold_fill = PatternFill(start_color="FFF8DC", end_color="FFF8DC", fill_type="solid") # Cornsilk Gold
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D4AF37'),
        right=Side(style='thin', color='D4AF37'),
        top=Side(style='thin', color='D4AF37'),
        bottom=Side(style='thin', color='D4AF37')
    )
    
    # -------------------------------------------------------------
    # SHEET 1: MASTER WEDDING SUMMARY
    # -------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Executive Summary")
    ws_summary.views.sheetView[0].showGridLines = True
    
    ws_summary.merge_cells("A1:D1")
    ws_summary["A1"] = f"💍 {wedding.get('title', 'Royal Wedding')} - Executive Catering & Headcount Summary"
    ws_summary["A1"].font = Font(name="Calibri", size=15, bold=True, color="E5C158")
    ws_summary["A1"].fill = PatternFill(start_color="0C1017", end_color="0C1017", fill_type="solid")
    ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 36
    
    summary_data = [
        ("Total Invited Relatives (Households)", stats.get("total_invited", 0), "Total primary invitations sent out via WhatsApp"),
        ("Confirmed Attendee Headcount", stats.get("total_headcount", 0), "Exact number of plates required for lunch/dinner catering"),
        ("RSVP Confirmed Households", stats.get("confirmed", 0), "Relatives who completed chatbot verification"),
        ("Pending Follow-ups", stats.get("pending", 0), "Awaiting relative response on WhatsApp"),
        ("Wedding Date", wedding.get("wedding_date", "TBD"), "Event Schedule"),
        ("City / Venue", wedding.get("wedding_city", "Gujarat"), "Location of Reception")
    ]
    
    ws_summary.append([])
    ws_summary.append(["Key Metric", "Value", "Catering & Event Impact"])
    ws_summary.row_dimensions[3].height = 24
    for col in range(1, 4):
        cell = ws_summary.cell(row=3, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    for row_idx, item in enumerate(summary_data, start=4):
        ws_summary.append([item[0], item[1], item[2]])
        ws_summary.row_dimensions[row_idx].height = 22
        ws_summary.cell(row=row_idx, column=1).font = bold_font
        ws_summary.cell(row=row_idx, column=2).font = bold_font
        ws_summary.cell(row=row_idx, column=2).alignment = center_align
        ws_summary.cell(row=row_idx, column=3).font = regular_font
        for c_idx in range(1, 4):
            ws_summary.cell(row=row_idx, column=c_idx).border = thin_border
            
    ws_summary.column_dimensions["A"].width = 38
    ws_summary.column_dimensions["B"].width = 24
    ws_summary.column_dimensions["C"].width = 45

    # -------------------------------------------------------------
    # SHEET 2: HOUSEHOLDS & CATERING LIST
    # -------------------------------------------------------------
    ws_households = wb.create_sheet(title="Households & Catering")
    ws_households.views.sheetView[0].showGridLines = True
    
    hh_headers = [
        "Family Head Name",
        "Mobile Number",
        "Side",
        "Invited By",
        "Category",
        "Total Guests Coming",
        "Spouse Name",
        "Children / Other Members",
        "RSVP Status",
        "Personal Guest Chat Link"
    ]
    ws_households.append(hh_headers)
    ws_households.row_dimensions[1].height = 26
    for col_num in range(1, len(hh_headers) + 1):
        cell = ws_households.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    for c_obj in contacts:
        head_name = c_obj.get("name", "")
        phone = c_obj.get("phone", "")
        side_val = c_obj.get("side", "")
        added_by = c_obj.get("added_by", "")
        cat = c_obj.get("category", "").title()
        headcount = c_obj.get("total_headcount", 1)
        spouse = c_obj.get("spouse_name", "") or "-"
        kids_members = []
        if c_obj.get("children_names"):
            kids_members.append(f"Kids: {c_obj.get('children_names')}")
        if c_obj.get("other_members"):
            kids_members.append(f"Others: {c_obj.get('other_members')}")
        members_str = "; ".join(kids_members) if kids_members else "-"
        status = c_obj.get("status", "WAITING")
        chat_link = f"{base_url.rstrip('/')}/chat/{c_obj.get('token', '')}"
        
        row_vals = [
            head_name,
            phone,
            side_val,
            added_by,
            cat,
            headcount,
            spouse,
            members_str,
            status,
            chat_link
        ]
        ws_households.append(row_vals)
        cur_row = ws_households.max_row
        ws_households.row_dimensions[cur_row].height = 20
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws_households.cell(row=cur_row, column=col_idx)
            cell.border = thin_border
            cell.font = regular_font
            if col_idx in [2, 3, 4, 5, 6, 9]:
                cell.alignment = center_align
            if col_idx == 6:
                cell.font = bold_font
                cell.fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
            if col_idx == 9 and status == "CONFIRMED":
                cell.font = bold_font
                cell.font = Font(name="Calibri", size=11, bold=True, color="137333")

    # Auto-adjust column widths
    for col in ws_households.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_households.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # -------------------------------------------------------------
    # SHEET 3: PENDING FOLLOW-UPS
    # -------------------------------------------------------------
    ws_pending = wb.create_sheet(title="Pending Follow-ups")
    ws_pending.views.sheetView[0].showGridLines = True
    
    pend_headers = [
        "Relative Name",
        "Mobile Number",
        "Side",
        "Category",
        "Status",
        "WhatsApp Direct Link"
    ]
    ws_pending.append(pend_headers)
    ws_pending.row_dimensions[1].height = 26
    for col_num in range(1, len(pend_headers) + 1):
        cell = ws_pending.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    pending_list = [c_obj for c_obj in contacts if c_obj.get("status") != "CONFIRMED"]
    for p_obj in pending_list:
        phone_num = p_obj.get("phone", "")
        clean_p = "".join(filter(str.isdigit, str(phone_num)))[-10:]
        wa_url = f"https://wa.me/91{clean_p}?text={wedding.get('groom_name','Groom')}%20Wedding%20RSVP:%20{base_url.rstrip('/')}/chat/{p_obj.get('token', '')}" if clean_p else "-"
        row_vals = [
            p_obj.get("name", ""),
            phone_num,
            p_obj.get("side", ""),
            p_obj.get("category", "").title(),
            p_obj.get("status", "WAITING"),
            wa_url
        ]
        ws_pending.append(row_vals)
        cur_row = ws_pending.max_row
        ws_pending.row_dimensions[cur_row].height = 20
        for col_idx in range(1, len(row_vals) + 1):
            cell = ws_pending.cell(row=cur_row, column=col_idx)
            cell.border = thin_border
            cell.font = regular_font
            if col_idx in [2, 3, 4, 5]:
                cell.alignment = center_align
                
    for col in ws_pending.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_pending.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(output_path)
    return output_path
